#!/usr/bin/env python3
"""Gera as planilhas (CSV + XLSX) e as tabelas em markdown do plano, e valida o plano.

Uso: python3 scripts/gerar.py
Falha (exit 1) se: algum dia não somar 4h técnicas + 1h de inglês; alguma semana tiver
menos de 50% de prática; algum tópico das seções 3/4 ficar sem semana; algum código de
tópico, trilha ou recurso não existir; alguma trilha tiver fora de 2-4 recursos principais.
"""
import csv
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from plano.candidatura import CHECKLIST, RESPOSTAS  # noqa: E402
from plano.cronograma import SEMANAS, ingles_da_semana  # noqa: E402
from plano.diagnostico import DIAGNOSTICO, GRUPO_PARA_DIAG  # noqa: E402
from plano.recursos import RECURSOS  # noqa: E402
from plano.topicos import TOPICOS  # noqa: E402
from plano.trilhas import TRILHAS  # noqa: E402

RAIZ = Path(__file__).resolve().parent.parent
PLANILHAS = RAIZ / "planilhas"
PLANO = RAIZ / "plano"
DIAG_DIR = RAIZ / "diagnostico"

DIAS = ["Seg", "Ter", "Qua", "Qui", "Sex", "Sáb"]
MODOS = {"E": "Estudo", "L": "Lab/exercício", "P": "Projeto", "W": "Escrita", "R": "Revisão de agente", "F": "Reforço flex"}
PRATICA = {"L", "P", "W", "R"}
STATUS = ["A fazer", "Em andamento", "Feito", "Pulado", "Adiado"]
NIVEIS = ["Iniciante", "Intermediário", "Avançado"]
CHECKPOINTS = [4, 8, 12, 16, 20, 24]

TOP = {t[0]: t for t in TOPICOS}
erros = []

# Marcação derivada: tópicos que um bloco exercita por natureza, além dos marcados à mão.
# (regex na atividade ou condição sobre trilha/modo/dia) -> códigos acrescentados.
REGRAS = [
    (lambda bl, d: bl["modo"] == "R", "VA01 VR08 LI05 LI07"),          # lab AI-native = delegar e verificar agentes
    (lambda bl, d: bl["trilha"] == "T05" and bl["modo"] in "PL", "VD03 VS12"),
    (lambda bl, d: d == 5 and bl["modo"] == "P", "EA01 EA02 EA03"),     # demo/release de sábado
    (lambda bl, d: bl["trilha"] == "T15" and re.search(r"[Jj]únior", bl["atividade"]), "VR06 LI01 LI02 VC02"),
    (lambda bl, d: re.search(r"Update semanal|Checkpoint", bl["atividade"]) and bl["trilha"] == "T15", "PC07 PC03"),
    (lambda bl, d: re.search(r"\bADR", bl["atividade"]), "OP07 EA13"),
    (lambda bl, d: re.search(r"[Dd]esign doc", bl["atividade"]), "VL01 EA06 PC01"),
    (lambda bl, d: re.search(r"[Pp]ostmortem", bl["atividade"]), "CQ10 PC02"),
    (lambda bl, d: re.search(r"[Rr]unbook", bl["atividade"]), "CD09"),
    (lambda bl, d: re.search(r"\bRelease\b", bl["atividade"]) and bl["modo"] == "P", "CD06 CQ03 V304"),
    (lambda bl, d: re.search(r"\bPRs?\b", bl["atividade"]), "LT05"),
    (lambda bl, d: re.search(r"Postgres|Cloud SQL", bl["atividade"]), "VS03 DA01"),
    (lambda bl, d: re.search(r"\bSQL\b", bl["atividade"]), "LT04"),
    (lambda bl, d: re.search(r"BigQuery", bl["atividade"]), "VS10 DA02 VB05"),
    (lambda bl, d: re.search(r"Next\.js", bl["atividade"]), "VS05 LT03"),
    (lambda bl, d: re.search(r"Stripe", bl["atividade"]), "VS15"),
    (lambda bl, d: re.search(r"Airtable", bl["atividade"]), "VS19"),
    (lambda bl, d: re.search(r"Pub/Sub", bl["atividade"]), "VS09"),
    (lambda bl, d: re.search(r"Cloud Run", bl["atividade"]), "VS07"),
    (lambda bl, d: re.search(r"Temporal|workflow|CaseLifecycle", bl["atividade"]), "VS12"),
    (lambda bl, d: re.search(r"\(Go\)|em Go\b", bl["atividade"]), "VS01 LT02"),
    (lambda bl, d: re.search(r"\(Python\)|em Python\b", bl["atividade"]), "VS02 LT01"),
]


def derivar(bl, d):
    extra = []
    for cond, codes in REGRAS:
        if cond(bl, d):
            extra += [c for c in codes.split() if c not in bl["topicos"] and c not in extra]
    return bl["topicos"] + extra


def fmt_h(h):
    return f"{h:g}".replace(".", ",")


def pct1(x):
    return f"{x:.1%}".replace(".", ",")


def nome_trilha(tid):
    return f"{tid} {TRILHAS[tid][0]}" if tid != "FLEX" else TRILHAS[tid][0]


def links(ids):
    return " ; ".join(f"{RECURSOS[r][2]} — {RECURSOS[r][3]}" for r in ids)


def topicos_txt(codes):
    return "; ".join(f"{c} {TOP[c][3]}" for c in codes)


# ------------------------------------------------------------------ montar linhas
linhas = []  # cronograma diário
cobertura = defaultdict(set)
blocos_por_topico = defaultdict(int)
uso_recurso = defaultdict(set)
horas_trilha = defaultdict(float)
pratica_trilha = defaultdict(float)
horas_semana_trilha = defaultdict(lambda: defaultdict(float))
resumo = []

for s in SEMANAS:
    n = s["n"]
    tec_total = pratica = 0.0
    trilhas_semana = defaultdict(float)
    en = ingles_da_semana(n)
    for d, blocos in enumerate(s["dias"]):
        soma = sum(bl["horas"] for bl in blocos)
        if abs(soma - 4) > 1e-9:
            erros.append(f"Semana {n} dia {d + 1}: {soma}h técnicas (esperado 4h)")
        for bl in blocos:
            if bl["trilha"] not in TRILHAS:
                erros.append(f"Semana {n}: trilha inexistente {bl['trilha']}")
            if bl["trilha"] != "FLEX":
                bl = {**bl, "topicos": derivar(bl, d)}
            for c in bl["topicos"]:
                if c not in TOP:
                    erros.append(f"Semana {n}: tópico inexistente {c}")
                cobertura[c].add(n)
                blocos_por_topico[c] += 1
            for r in bl["recursos"]:
                if r not in RECURSOS:
                    erros.append(f"Semana {n}: recurso inexistente {r}")
                uso_recurso[r].add(n)
            tec_total += bl["horas"]
            horas_trilha[bl["trilha"]] += bl["horas"]
            trilhas_semana[bl["trilha"]] += bl["horas"]
            horas_semana_trilha[bl["trilha"]][n] += bl["horas"]
            if bl["modo"] in PRATICA:
                pratica += bl["horas"]
                pratica_trilha[bl["trilha"]] += bl["horas"]
            linhas.append({
                "Semana": n, "Dia": f"{d + 1} ({DIAS[d]})", "Bloco": "Técnico",
                "Trilha": nome_trilha(bl["trilha"]),
                "Tópico da vaga coberto": topicos_txt(bl["topicos"]) if bl["topicos"] else "Definido pelo diagnóstico",
                "Atividade concreta": bl["atividade"],
                "Recurso (link)": links(bl["recursos"]) if bl["recursos"] else "Projeto LeaveFlow / materiais do repositório",
                "Horas": bl["horas"], "Entregável": bl["entregavel"], "Status": "A fazer",
                "Modo": MODOS[bl["modo"]],
            })
        # bloco de inglês do dia
        tcodes, ativ, recs, entreg = en[d]
        tcodes = tcodes.split()
        recs = [r for r in recs.split(",") if r]
        for c in tcodes:
            if c not in TOP:
                erros.append(f"Semana {n} inglês: tópico inexistente {c}")
            cobertura[c].add(n)
            blocos_por_topico[c] += 1
        for r in recs:
            if r not in RECURSOS:
                erros.append(f"Semana {n} inglês: recurso inexistente {r}")
            uso_recurso[r].add(n)
        horas_trilha["T17"] += 1
        horas_semana_trilha["T17"][n] += 1
        linhas.append({
            "Semana": n, "Dia": f"{d + 1} ({DIAS[d]})", "Bloco": "Inglês",
            "Trilha": nome_trilha("T17"), "Tópico da vaga coberto": topicos_txt(tcodes),
            "Atividade concreta": ativ,
            "Recurso (link)": links(recs) if recs else "Gravação própria (celular/webcam) + rubrica do item F",
            "Horas": 1, "Entregável": entreg, "Status": "A fazer", "Modo": "Inglês",
        })
    pct = pratica / tec_total
    if pct < 0.5:
        erros.append(f"Semana {n}: prática {pct:.0%} < 50%")
    resumo.append({
        "Semana": n, "Fase": s["fase"], "Tema": s["tema"], "Objetivos": s["objetivos"],
        "Marco do projeto": s["marco"], "Critério de pronto": s["pronto"],
        "Artefatos de liderança": s["lideranca"],
        "Checkpoint de revisão": "Sim: Matriz + horas reais + replanejamento" if s["checkpoint"] else "—",
        "Horas técnicas": tec_total, "Horas de inglês": 6, "% prática (técnico)": f"{pct:.0%}",
        "Horas por trilha": "; ".join(f"{k} {fmt_h(v)}h" for k, v in sorted(trilhas_semana.items())),
    })

faltando = [t[0] for t in TOPICOS if t[0] not in cobertura]
if faltando:
    erros.append("Tópicos sem semana: " + ", ".join(f"{c} ({TOP[c][3]})" for c in faltando))

for tid in TRILHAS:
    if tid in ("T00", "FLEX"):
        continue
    np = sum(1 for r in RECURSOS.values() if r[0] == tid and r[1] == "Principal")
    if not 2 <= np <= 4:
        erros.append(f"Trilha {tid}: {np} recursos principais (esperado 2-4)")

for r, v in RECURSOS.items():
    if v[0] not in TRILHAS:
        erros.append(f"Recurso {r}: trilha inexistente {v[0]}")

if erros:
    print("ERROS DE VALIDAÇÃO:")
    for e in erros:
        print(" -", e)
    sys.exit(1)

# ------------------------------------------------------------------ tabelas
COL_CRON = ["Semana", "Dia", "Bloco", "Trilha", "Tópico da vaga coberto", "Atividade concreta",
            "Recurso (link)", "Horas", "Entregável", "Status", "Modo"]

diag_rows = [{
    "ID": k, "Grupo": v[0], "Exercício": v[1],
    "Critérios de nível": f"Iniciante: {v[2]}\nIntermediário: {v[3]}\nAvançado: {v[4]}",
    "Tempo": f"{v[5]} min", "Resultado (eu preencho)": "", "Nível": "",
    "Trilhas afetadas": v[6],
} for k, v in DIAGNOSTICO.items()]
COL_DIAG = list(diag_rows[0].keys())

rec_rows = [{
    "ID": k, "Trilha": nome_trilha(v[0]), "Papel": v[1], "Nome": v[2], "Link": v[3], "Tipo": v[4],
    "Custo": v[5], "Horas estimadas": v[6], "Nível": v[7],
    "Usado nas semanas": ", ".join(str(x) for x in sorted(uso_recurso[k])) or "Consulta livre",
} for k, v in RECURSOS.items()]
COL_REC = list(rec_rows[0].keys())


def diag_do_topico(t):
    if t[6] != "—":
        return t[6]
    return "— (não mensurável por exercício)"


mat_rows = []
for t in TOPICOS:
    row = {
        "ID": t[0], "Seção": t[1], "Grupo": t[2], "Tópico": t[3], "Trilha": nome_trilha(t[4]),
        "Prioridade": t[5], "Exercício de diagnóstico": diag_do_topico(t),
        "Nível no diagnóstico (preencher)": "",
        "Semanas em que é coberto": ", ".join(str(x) for x in sorted(cobertura[t[0]])),
        "Nº de blocos": blocos_por_topico[t[0]],
        "Como será comprovado": t[7], "Nível atual": "",
    }
    for cp in CHECKPOINTS:
        row[f"CP sem. {cp}"] = ""
    mat_rows.append(row)
COL_MAT = list(mat_rows[0].keys())

cand_rows = [{
    "Tipo": "Pergunta", "Item": q, "Rascunho (EN)": rasc, "Evidência a usar": ev,
    "Cuidados / o que preencher": cuid, "Finalizar em": quando, "Status": "Rascunho",
} for q, rasc, ev, cuid, quando in RESPOSTAS]
cand_rows += [{
    "Tipo": "Checklist antes de aplicar", "Item": item, "Rascunho (EN)": "", "Evidência a usar": "",
    "Cuidados / o que preencher": "", "Finalizar em": "Semana 24", "Status": "A fazer",
} for item in CHECKLIST]
COL_CAND = list(cand_rows[0].keys())

tot_tec = sum(v for k, v in horas_trilha.items() if k != "T17")
horas_rows = []
for tid, tr in TRILHAS.items():
    h = horas_trilha.get(tid, 0)
    pr = pratica_trilha.get(tid, 0)
    horas_rows.append({
        "Trilha": nome_trilha(tid), "Prioridade": tr[3], "Nível-alvo": tr[2],
        "Horas no plano": h,
        "% do tempo técnico": "fora das 4h" if tid == "T17" else pct1(h / tot_tec),
        "Horas de prática": "—" if tid in ("T17", "FLEX") else pr,
        "% prática na trilha": "—" if tid in ("T17", "FLEX") or not h else f"{pr / h:.0%}",
    })
COL_HORAS = list(horas_rows[0].keys())

PLANILHAS.mkdir(exist_ok=True)


def write_csv(nome, cols, rows):
    with open(PLANILHAS / nome, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        for r in rows:
            w.writerow({c: (fmt_h(r[c]) if isinstance(r[c], float) else r[c]) for c in cols})


write_csv("01-diagnostico.csv", COL_DIAG, diag_rows)
write_csv("02-cronograma-diario.csv", COL_CRON, linhas)
write_csv("03-resumo-semanal.csv", list(resumo[0].keys()), resumo)
write_csv("04-recursos.csv", COL_REC, rec_rows)
write_csv("05-matriz-de-cobertura.csv", COL_MAT, mat_rows)
write_csv("06-candidatura.csv", COL_CAND, cand_rows)
write_csv("07-horas-por-trilha.csv", COL_HORAS, horas_rows)

# ------------------------------------------------------------------ XLSX
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
EN_FILL = PatternFill("solid", fgColor="E8F1FB")
FLEX_FILL = PatternFill("solid", fgColor="FFF6D6")
CP_FILL = PatternFill("solid", fgColor="E9F7EF")
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()
ws0 = wb.active
ws0.title = "Leia-me"
leia = [
    ["Plano de estudos: Engineering Lead (Backend Systems & Growth) — 24 semanas"],
    [""],
    ["Abas"],
    ["Diagnóstico", "17 exercícios da semana 1 (e 2) com critérios objetivos. Preencha Resultado e Nível."],
    ["Cronograma diário", "144 dias × (blocos técnicos de 4h + 1h de inglês). Atualize a coluna Status todo dia."],
    ["Resumo semanal", "Objetivos, marco do projeto, critério de pronto e checkpoints (semanas 4, 8, 12, 16, 20, 24)."],
    ["Recursos", "Banco de recursos com links conferidos (Principal = item B; Apoio = docs usadas em atividades)."],
    ["Matriz de cobertura", "Todos os tópicos das seções 3 e 4, um por linha, com semanas calculadas a partir do cronograma."],
    ["Candidatura", "Rascunhos das 7 respostas (em inglês) + checklist antes de aplicar."],
    ["Horas por trilha", "Distribuição de horas e percentual de prática por trilha."],
    [""],
    ["Números do plano"],
    ["Horas técnicas", fmt_h(tot_tec)],
    ["Horas de inglês", fmt_h(horas_trilha['T17'])],
    ["% prática (técnico)", f"{sum(pratica_trilha.values()) / tot_tec:.0%} (reforço flex conta como 0%)"],
    ["Tópicos na matriz", str(len(TOPICOS))],
    [""],
    ["Gerado por scripts/gerar.py. Edite os dados em scripts/plano/*.py e rode de novo para regenerar."],
]
for r in leia:
    ws0.append(r)
ws0["A1"].font = Font(bold=True, size=14)
for c in ("A3", "A12"):
    ws0[c].font = Font(bold=True)
ws0.column_dimensions["A"].width = 24
ws0.column_dimensions["B"].width = 110


def sheet(titulo, cols, rows, larguras, fill_fn=None, validacoes=None, link_col=None):
    ws = wb.create_sheet(titulo)
    ws.append(cols)
    for c in ws[1]:
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, WRAP
    for r in rows:
        ws.append([r[c] for c in cols])
    for i, w in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        fill = fill_fn(row) if fill_fn else None
        for c in row:
            c.alignment = WRAP
            if fill:
                c.fill = fill
    if link_col:
        idx = cols.index(link_col) + 1
        for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("http") and " " not in c.value:
                    c.hyperlink = c.value
                    c.font = Font(color="0563C1", underline="single")
    for col_nome, opcoes in (validacoes or {}).items():
        idx = cols.index(col_nome) + 1
        letra = get_column_letter(idx)
        dv = DataValidation(type="list", formula1='"' + ",".join(opcoes) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{letra}2:{letra}{len(rows) + 1}")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


sheet("Diagnóstico", COL_DIAG, diag_rows, [6, 30, 70, 70, 9, 40, 16, 12],
      validacoes={"Nível": NIVEIS})


def fill_cron(row):
    if row[2].value == "Inglês":
        return EN_FILL
    if row[10].value == "Reforço flex":
        return FLEX_FILL
    return None


sheet("Cronograma diário", COL_CRON, linhas, [8, 9, 9, 26, 40, 70, 45, 7, 32, 13, 14],
      fill_fn=fill_cron, validacoes={"Status": STATUS})
sheet("Resumo semanal", list(resumo[0].keys()), resumo, [8, 22, 30, 45, 40, 45, 35, 22, 9, 9, 10, 45],
      fill_fn=lambda row: CP_FILL if row[7].value != "—" else None)
sheet("Recursos", COL_REC, rec_rows, [7, 26, 10, 45, 50, 18, 22, 9, 18, 18], link_col="Link")
sheet("Matriz de cobertura", COL_MAT, mat_rows,
      [6, 16, 24, 45, 26, 12, 14, 14, 22, 8, 45, 12] + [10] * len(CHECKPOINTS),
      validacoes={"Nível no diagnóstico (preencher)": NIVEIS, "Nível atual": NIVEIS,
                  **{f"CP sem. {cp}": NIVEIS for cp in CHECKPOINTS}})
sheet("Candidatura", COL_CAND, cand_rows, [14, 40, 90, 35, 35, 18, 12],
      validacoes={"Status": ["Rascunho", "Revisado", "Final", "A fazer", "Feito"]})
sheet("Horas por trilha", COL_HORAS, horas_rows, [55, 22, 40, 10, 12, 10, 12])
wb.save(PLANILHAS / "plano-de-estudos.xlsx")


# ------------------------------------------------------------------ Markdown gerado
def md_table(cols, rows):
    out = ["| " + " | ".join(cols) + " |", "|" + "|".join(["---"] * len(cols)) + "|"]
    for r in rows:
        out.append("| " + " | ".join(str(r[c]).replace("|", "/").replace("\n", "<br>") for c in cols) + " |")
    return "\n".join(out)


AVISO = "<!-- Arquivo gerado por scripts/gerar.py — não edite à mão. -->\n\n"

# B — banco de recursos
partes = [AVISO + "# B) Banco de recursos\n",
          "Links conferidos por busca na web em outubro de 2026. **Principal** = os 2 a 4 materiais de cada trilha "
          "(item B do pedido). **Apoio** = documentação usada em atividades específicas do cronograma. "
          "Horas são estimativas de consumo do material (não o tempo de prática no projeto). "
          "\"Google Skills\" é o novo nome do Google Cloud Skills Boost (os links antigos redirecionam).\n",
          "A trilha T00 (diagnóstico) usa os materiais da pasta [`diagnostico/`](../diagnostico/README.md); "
          "a de reforço flex usa os recursos da trilha que estiver sendo reforçada.\n"]
for tid, tr in TRILHAS.items():
    regs = [(k, v) for k, v in RECURSOS.items() if v[0] == tid]
    if not regs:
        continue
    partes.append(f"\n## {tid} — {tr[0]}\n")
    cols = ["ID", "Papel", "Material", "Tipo", "Custo", "Horas", "Nível", "Semanas"]
    rows = [{"ID": k, "Papel": v[1], "Material": f"[{v[2]}]({v[3]})", "Tipo": v[4], "Custo": v[5],
             "Horas": fmt_h(float(v[6])), "Nível": v[7],
             "Semanas": ", ".join(str(x) for x in sorted(uso_recurso[k])) or "consulta"}
            for k, v in sorted(regs, key=lambda kv: (kv[1][1] != "Principal", kv[0]))]
    partes.append(md_table(cols, rows))
(PLANO / "B-banco-de-recursos.md").write_text("\n".join(partes) + "\n", encoding="utf-8")

# Matriz de cobertura
cols = ["ID", "Grupo", "Tópico", "Trilha", "Prioridade", "Diagnóstico", "Semanas", "Como será comprovado"]
rows = [{"ID": r["ID"], "Grupo": r["Grupo"], "Tópico": r["Tópico"], "Trilha": r["Trilha"].split(" ")[0],
         "Prioridade": r["Prioridade"], "Diagnóstico": r["Exercício de diagnóstico"],
         "Semanas": r["Semanas em que é coberto"], "Como será comprovado": r["Como será comprovado"]}
        for r in mat_rows]
n4 = sum(1 for t in TOPICOS if t[1].startswith("Seção 4"))
n3 = len(TOPICOS) - n4
txt = (AVISO + "# Matriz de cobertura\n\n"
       f"Todos os tópicos das seções 3 e 4 do pedido, um por linha: **{n4} itens da seção 4** "
       f"(o pedido fala em 81, mas a lista colada tem {n4}; todos foram incluídos) e **{n3} da seção 3** "
       "(responsabilidades, padrões, Airtable, Growth, liderança, AI-native, 30-60-90, requisitos, diferenciais, "
       "cultura, stack e as 7 perguntas). As semanas são calculadas a partir do cronograma; o gerador falha se algum "
       "tópico ficar sem semana. Os níveis (diagnóstico, atual e por checkpoint) são preenchidos na planilha "
       "[`05-matriz-de-cobertura.csv`](../planilhas/05-matriz-de-cobertura.csv) ou na aba do XLSX.\n\n"
       + md_table(cols, rows) + "\n")
(PLANO / "matriz-de-cobertura.md").write_text(txt, encoding="utf-8")

t15_regular = sum(ln["Horas"] for ln in linhas
                  if ln["Bloco"] == "Técnico" and ln["Trilha"].startswith("T15 ") and 3 <= ln["Semana"] <= 22)

# G — resumo semanal + horas
cols_r = ["Semana", "Tema", "Marco do projeto", "Critério de pronto", "Artefatos de liderança", "Checkpoint de revisão", "% prática (técnico)"]
pct_total = sum(pratica_trilha.values()) / tot_tec
txt = (AVISO + "# G) Cronograma em formato de planilha\n\n"
       "## Arquivos\n\n"
       "| Aba pedida | CSV (colar no Sheets/Excel) | Linhas |\n|---|---|---|\n"
       f"| Diagnóstico | [`01-diagnostico.csv`](../planilhas/01-diagnostico.csv) | {len(diag_rows)} |\n"
       f"| Cronograma diário | [`02-cronograma-diario.csv`](../planilhas/02-cronograma-diario.csv) | {len(linhas)} |\n"
       f"| Resumo semanal | [`03-resumo-semanal.csv`](../planilhas/03-resumo-semanal.csv) | {len(resumo)} |\n"
       f"| Recursos | [`04-recursos.csv`](../planilhas/04-recursos.csv) | {len(rec_rows)} |\n"
       f"| Matriz de cobertura | [`05-matriz-de-cobertura.csv`](../planilhas/05-matriz-de-cobertura.csv) | {len(mat_rows)} |\n"
       f"| Candidatura | [`06-candidatura.csv`](../planilhas/06-candidatura.csv) | {len(cand_rows)} |\n"
       f"| (extra) Horas por trilha | [`07-horas-por-trilha.csv`](../planilhas/07-horas-por-trilha.csv) | {len(horas_rows)} |\n\n"
       "Tudo junto, com abas, filtros, listas suspensas de Status/Nível e links clicáveis: "
       "[`planilhas/plano-de-estudos.xlsx`](../planilhas/plano-de-estudos.xlsx) (abre no Excel e no Google Sheets: "
       "Arquivo → Importar → Fazer upload).\n\n"
       "**Para colar um CSV no Google Sheets:** Arquivo → Importar → Fazer upload → escolha o CSV → "
       "\"Substituir planilha atual\" ou \"Inserir nova(s) página(s)\"; separador: vírgula. No Excel: Dados → De Texto/CSV → "
       "codificação UTF-8. Os CSVs usam UTF-8 com BOM, então acentos abrem corretamente.\n\n"
       "A coluna **Modo** (extra, no fim do cronograma) separa Estudo, Lab/exercício, Projeto, Escrita, Revisão de agente "
       "e Reforço flex; é ela que prova a regra de ≥50% de prática.\n\n"
       "## Números do plano\n\n"
       f"- {len(SEMANAS)} semanas × 6 dias = {len(SEMANAS) * 6} dias; {fmt_h(tot_tec)} h técnicas + "
       f"{fmt_h(horas_trilha['T17'])} h de inglês.\n"
       f"- Prática no tempo técnico: **{pct_total:.0%}** (laboratórios, projeto, escrita de artefatos e revisão de agentes; "
       "o reforço flex conta como 0%). Nenhuma semana fica abaixo de 50%.\n"
       f"- Liderança e comunicação (T15): {fmt_h(horas_trilha['T15'])} h = "
       f"{pct1(horas_trilha['T15'] / tot_tec)} do tempo técnico (alvo do pedido: ~10%; {pct1(t15_regular / (20 * 24))} nas semanas 3-22; "
       "as semanas 1, 2, 23 e 24 concentram diagnóstico, planejamento e avaliação de lacunas).\n"
       f"- Tópicos na matriz: {len(TOPICOS)}, todos com pelo menos uma semana.\n"
       "- **Horas por trilha contam a trilha principal de cada bloco.** Padrões de T02-T04 (SQL, idempotência, outbox, "
       "retries) também são praticados dentro de blocos de T05, T08, T09 e T10; a coluna 'Nº de blocos' da matriz mostra "
       "quantas vezes cada tópico aparece.\n\n"
       "## Horas por trilha\n\n" + md_table(COL_HORAS, [{**r, "Horas no plano": fmt_h(r["Horas no plano"]),
                                                          "Horas de prática": r["Horas de prática"] if isinstance(r["Horas de prática"], str) else fmt_h(r["Horas de prática"])}
                                                         for r in horas_rows]) +
       "\n\n## Resumo semanal\n\n" + md_table(cols_r, resumo) + "\n")
(PLANO / "G-cronograma.md").write_text(txt, encoding="utf-8")

# Diagnóstico — tabela de exercícios
cols_d = ["ID", "Grupo", "Exercício", "Iniciante", "Intermediário", "Avançado", "Tempo"]
rows_d = [{"ID": k, "Grupo": v[0], "Exercício": v[1], "Iniciante": v[2], "Intermediário": v[3],
           "Avançado": v[4], "Tempo": f"{v[5]} min"} for k, v in DIAGNOSTICO.items()]
mapa = "\n".join(f"| {g} | {d} |" for g, d in GRUPO_PARA_DIAG.items())
txt = (AVISO + "# Exercícios de diagnóstico e critérios de nível\n\n"
       "Materiais e instruções gerais: [`README.md`](README.md). Registre resultado e nível em "
       "[`planilhas/01-diagnostico.csv`](../planilhas/01-diagnostico.csv) (ou na aba Diagnóstico do XLSX).\n\n"
       "## Grupo da seção 4 → exercícios\n\n| Grupo | Exercícios |\n|---|---|\n" + mapa +
       "\n| (seção 3) AI-native | D15 |\n| (seção 3) Growth/frontend | D16 |\n\n## Exercícios\n\n" + md_table(cols_d, rows_d) + "\n")
(DIAG_DIR / "exercicios.md").write_text(txt, encoding="utf-8")

# ------------------------------------------------------------------ Documentos a partir de templates
TEMPLATES = Path(__file__).resolve().parent / "plano" / "templates"
modo_h = defaultdict(float)
for ln in linhas:
    if ln["Bloco"] == "Técnico":
        modo_h[ln["Modo"]] += ln["Horas"]


def pct(h):
    return f"{h / tot_tec:.0%}"


def semanas_foco(tid):
    ws = [n for n, h in sorted(horas_semana_trilha[tid].items()) if h >= 3]
    if len(ws) == len(SEMANAS):
        return "todas"
    return ", ".join(map(str, ws)) or "contínua (blocos curtos)"


tab_trilhas = md_table(
    ["Trilha", "Objetivo", "Nível-alvo", "Prioridade", "Depende de", "Horas", "Semanas com foco (≥3 h)", "Se atrasar, pule"],
    [{"Trilha": f"**{tid}** {tr[0]}", "Objetivo": tr[1], "Nível-alvo": tr[2], "Prioridade": tr[3], "Depende de": tr[4],
      "Horas": fmt_h(horas_trilha.get(tid, 0)), "Semanas com foco (≥3 h)": semanas_foco(tid), "Se atrasar, pule": tr[5]}
     for tid, tr in TRILHAS.items()])
tab_marcos = md_table(["Semana", "Tema", "Marco do projeto", "Critério de pronto", "Artefatos de liderança"], resumo)
labs = [ln for ln in linhas if ln["Modo"] == "Revisão de agente"]
tab_labs = md_table(["Semana", "Dia", "Atividade concreta", "Entregável"], labs)
lid = [ln for ln in linhas if ln["Trilha"].startswith("T15 ")]
tab_lid = md_table(["Semana", "Dia", "Horas", "Atividade concreta", "Entregável"],
                   [{**ln, "Horas": fmt_h(float(ln["Horas"]))} for ln in lid])
from plano.cronograma import INGLES  # noqa: E402
ing_rows = [{"Semana": 1, "Vocabulário": "—", "STAR (quarta)": "—", "Documento (quinta)": "Amostra escrita do D17",
             "Conversa/mock (sexta)": "Resposta não preparada à Q1", "Vídeo (sábado)": "Diagnóstico D17: EF SET + vídeo de 2 min"}]
for n, (voc, q, doc, mock, video) in sorted(INGLES.items()):
    ing_rows.append({"Semana": n, "Vocabulário": voc, "STAR (quarta)": TOP[q][3].split(" ")[0],
                     "Documento (quinta)": doc, "Conversa/mock (sexta)": mock, "Vídeo (sábado)": video})
tab_ing = md_table(list(ing_rows[0].keys()), ing_rows)

valores = {
    "HORAS_TEC": fmt_h(tot_tec), "HORAS_EN": fmt_h(horas_trilha["T17"]), "PCT_PRATICA": f"{pct_total:.0%}",
    "PCT_PROJETO": pct(modo_h["Projeto"]), "PCT_ESCRITA": pct(modo_h["Escrita"]), "PCT_ESTUDO": pct(modo_h["Estudo"]),
    "PCT_LAB": pct(modo_h["Lab/exercício"]), "PCT_REVISAO": pct(modo_h["Revisão de agente"]),
    "PCT_FLEX": pct(modo_h["Reforço flex"]),
    "T15_HORAS": fmt_h(horas_trilha["T15"]), "T15_PCT": pct1(horas_trilha["T15"] / tot_tec),
    "T15_PCT_REGULAR": pct1(t15_regular / (20 * 24)),
    "N_TOPICOS": str(len(TOPICOS)), "N_S4": str(n4), "N_S3": str(n3), "N_RECURSOS": str(len(RECURSOS)),
    "N_TRILHAS": str(len([t for t in TRILHAS if t != "FLEX"])), "N_LABS": str(len(labs)),
    "TABELA_TRILHAS": tab_trilhas, "TABELA_MARCOS": tab_marcos, "TABELA_LABS": tab_labs,
    "TABELA_LIDERANCA": tab_lid, "TABELA_INGLES": tab_ing,
}
SAIDAS = {
    "README.md": RAIZ / "README.md",
    "00-visao-geral-e-ajuste.md": PLANO / "00-visao-geral-e-ajuste.md",
    "A-mapa-de-trilhas.md": PLANO / "A-mapa-de-trilhas.md",
    "C-projeto-portfolio.md": PLANO / "C-projeto-portfolio.md",
    "D-fluxo-ai-native.md": PLANO / "D-fluxo-ai-native.md",
    "E-lideranca.md": PLANO / "E-lideranca.md",
    "F-ingles.md": PLANO / "F-ingles.md",
}
for nome, destino in SAIDAS.items():
    txt = (TEMPLATES / nome).read_text(encoding="utf-8")
    for k, v in valores.items():
        txt = txt.replace("{{" + k + "}}", v)
    faltam = re.findall(r"\{\{[A-Z_0-9]+\}\}", txt)
    if faltam:
        print(f"ERRO: placeholders sem valor em {nome}: {faltam}")
        sys.exit(1)
    destino.write_text(AVISO.replace("não edite à mão", f"edite scripts/plano/templates/{nome}") + txt, encoding="utf-8")

# Candidatura em markdown
partes = [AVISO + "# Candidatura: rascunhos das 7 respostas e checklist\n",
          "Rascunhos em inglês (a candidatura é em inglês). Trechos entre colchetes são dados que só você tem: "
          "preencha com números reais ou apague. **Nunca apresente o LeaveFlow como sistema de produção com usuários "
          "reais.** Versões finais na semana 24; cada pergunta é praticada em voz alta 3 vezes antes (trilha F). "
          "A mesma tabela está em [`planilhas/06-candidatura.csv`](../planilhas/06-candidatura.csv).\n"]
for q, rasc, ev, cuid, quando in RESPOSTAS:
    partes.append(f"\n## {q}\n\n{rasc}\n\n- **Evidência a usar:** {ev}\n- **Cuidados:** {cuid}\n- **Finalizar:** {quando}\n")
partes.append("\n## Checklist antes de se candidatar\n")
partes += [f"- [ ] {item}" for item in CHECKLIST]
(PLANO / "candidatura.md").write_text("\n".join(partes) + "\n", encoding="utf-8")

print(f"OK: {len(SEMANAS)} semanas, {len(linhas)} linhas no cronograma, {len(TOPICOS)} tópicos cobertos, "
      f"{len(RECURSOS)} recursos, {fmt_h(tot_tec)} h técnicas ({pct_total:.0%} prática), "
      f"T15 = {pct1(horas_trilha['T15'] / tot_tec)}.")
for tid in TRILHAS:
    print(f"  {tid:5} {fmt_h(horas_trilha.get(tid, 0)):>6} h")
